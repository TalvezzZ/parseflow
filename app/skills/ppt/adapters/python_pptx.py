import asyncio
import tempfile
from pathlib import Path

from app.documents.models import DocumentBlock, DocumentPage, DocumentResult, ParseContext, ProviderError
from app.skills.providers import Provider, ProviderManifest, ProviderResult


class PythonPptxProvider(Provider):
    """使用 python-pptx 提取页级文本、表格、图片和组合形状。"""

    manifest = ProviderManifest(
        name="ppt.normal.python-pptx",
        version="0.4.0",
        kind="presentation_parser",
        modes=["normal"],
        capabilities=["slides", "text", "tables", "images", "group_shapes", "markdown"],
    )

    async def parse(self, context: ParseContext) -> ProviderResult:
        return await asyncio.to_thread(self._parse_sync, context)

    def _parse_sync(self, context: ParseContext) -> ProviderResult:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError:
            return ProviderResult(status="failed", provider_name=self.manifest.name,
                                  error=ProviderError(code="provider_unavailable", message="未安装 python-pptx", retryable=False))
        try:
            presentation = Presentation(context.file.path)
        except Exception as exc:
            return ProviderResult(status="failed", provider_name=self.manifest.name,
                                  error=ProviderError(code="provider_failed", message=str(exc), retryable=False))
        artifact_dir = Path(str(context.metadata.get("artifact_dir") or tempfile.mkdtemp(prefix=f"parse-agent-{context.file.file_id}-pptx-")))
        images_dir = artifact_dir / "images"
        pages: list[DocumentPage] = []
        blocks: list[DocumentBlock] = []
        tables: list[dict] = []
        images: list[dict] = []
        markdown: list[str] = []

        def visit(shape, slide_number: int, slide_blocks: list[DocumentBlock], slide_lines: list[str]) -> None:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                for child in shape.shapes:
                    visit(child, slide_number, slide_blocks, slide_lines)
                return
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                table = {"slide_number": slide_number, "shape_id": shape.shape_id, "rows": rows}
                tables.append(table)
                text = "\n".join(" | ".join(row) for row in rows)
                slide_blocks.append(DocumentBlock(kind="table", text=text, page_number=slide_number, metadata=table))
                slide_lines.append(text)
                return
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images_dir.mkdir(parents=True, exist_ok=True)
                suffix = Path(getattr(shape.image, "filename", "image.png")).suffix or ".png"
                path = images_dir / f"slide-{slide_number:04d}-shape-{shape.shape_id}{suffix}"
                path.write_bytes(shape.image.blob)
                item = {"slide_number": slide_number, "shape_id": shape.shape_id, "filename": path.name, "path": str(path), "relative_path": f"images/{path.name}"}
                images.append(item)
                slide_blocks.append(DocumentBlock(kind="image", page_number=slide_number, metadata=item))
                return
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    kind = "heading" if getattr(shape, "is_placeholder", False) and getattr(shape.placeholder_format, "type", None) else "text"
                    slide_blocks.append(DocumentBlock(kind=kind, text=text, page_number=slide_number, metadata={"shape_id": shape.shape_id}))
                    slide_lines.append(text)

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_blocks: list[DocumentBlock] = []
            slide_lines: list[str] = []
            for shape in slide.shapes:
                visit(shape, slide_number, slide_blocks, slide_lines)
            pages.append(DocumentPage(page_number=slide_number, text="\n".join(slide_lines), blocks=slide_blocks))
            blocks.extend(slide_blocks)
            markdown.append(f"## Slide {slide_number}\n\n" + "\n\n".join(slide_lines))

        if not pages:
            return ProviderResult(status="failed", provider_name=self.manifest.name,
                                  error=ProviderError(code="quality_insufficient", message="PPT 未包含幻灯片", retryable=False))
        document = DocumentResult(document_id=context.file.file_id, source_file=context.file, document_type="pptx",
            pages=pages, blocks=blocks, tables=tables, images=images,
            representations={"plain_text": "\n\n".join(page.text for page in pages), "markdown": "\n\n".join(markdown)},
            metadata={"artifact_dir": str(artifact_dir), "slide_count": len(pages), "slide_width": presentation.slide_width, "slide_height": presentation.slide_height},
            parser={"name": self.manifest.name, "version": self.manifest.version},
            quality={"slides": len(pages), "text_blocks": sum(1 for block in blocks if block.kind in {"text", "heading"}), "tables": len(tables), "images": len(images)},
            extensions={"presentation": {"slide_count": len(pages)}})
        return ProviderResult(status="success", provider_name=self.manifest.name, document=document, metrics=document.quality)
