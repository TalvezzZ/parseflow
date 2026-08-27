from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.tasks.exports import ExportLimitExceeded, generate_document_exports


def test_generates_document_and_table_exports(tmp_path: Path) -> None:
    document = {"document_type": "csv", "representations": {"plain_text": "A B", "markdown": "# Data"},
                "tables": [{"name": "data", "rows": [["A", "B"], [1, "x|y"], [2, "=HYPERLINK(\"https://invalid\")"]]}]}
    paths = generate_document_exports(document, tmp_path, max_bytes=1024 * 1024)
    names = {path.name for path in paths}
    assert names == {"document.txt", "document.md", "document.json", "tables.xlsx", "table-001.csv", "table-001.json", "table-001.md"}
    assert json.loads((tmp_path / "exports/document.json").read_text())["document_type"] == "csv"
    with (tmp_path / "exports/table-001.csv").open(encoding="utf-8-sig") as handle:
        assert list(csv.reader(handle)) == [["A", "B"], ["1", "x|y"], ["2", "'=HYPERLINK(\"https://invalid\")"]]
    assert "x\\|y" in (tmp_path / "exports/table-001.md").read_text()
    formula_cell = load_workbook(tmp_path / "exports/tables.xlsx", read_only=True).active["B3"]
    assert formula_cell.data_type == "s" and formula_cell.value.startswith("'=")


def test_export_budget_rejects_oversized_output_without_partial_artifacts(tmp_path: Path) -> None:
    with pytest.raises(ExportLimitExceeded):
        generate_document_exports({"representations": {"plain_text": "x" * 8, "markdown": "y" * 8}}, tmp_path, max_bytes=10)
    assert not (tmp_path / "exports").exists()
    assert not (tmp_path / ".exports.tmp").exists()


def test_xlsx_sanitizes_invalid_and_case_duplicate_sheet_names(tmp_path: Path) -> None:
    document = {"tables": [{"name": "A/B", "rows": [[1]]}, {"name": "a_b", "rows": [[2]]}]}
    generate_document_exports(document, tmp_path, max_bytes=1024 * 1024)
    assert load_workbook(tmp_path / "exports/tables.xlsx", read_only=True).sheetnames == ["A_B", "a_b-2"]
