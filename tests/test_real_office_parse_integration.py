"""Real Office parsing through the public v0.8.0 upload/task API."""
from __future__ import annotations

import asyncio
import shutil
import subprocess
from pathlib import Path

import httpx
import pytest
from app.main import app


async def parse_upload(path: Path, content_type: str) -> dict:
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://testserver") as client:
        created = await client.post("/api/v1/tasks/parse", files={"file": (path.name, path.read_bytes(), content_type)})
        assert created.status_code == 202, created.text
        for _ in range(500):
            task = (await client.get(f"/api/v1/tasks/{created.json()['task_id']}")).json()
            if task["status"] in {"succeeded", "partial", "failed", "cancelled", "interrupted"}:
                return task
            await asyncio.sleep(.02)
    raise AssertionError("task did not finish")


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Sales"
    sheet.append(["Product", "Amount"]); sheet.append(["A", 2]); workbook.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    presentation = Presentation(); slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Quarterly Results"
    slide.shapes.add_textbox(1000000, 1500000, 5000000, 800000).text = "Revenue increased"
    presentation.save(path)


@pytest.mark.asyncio
async def test_real_xlsx_parse_upload_task(tmp_path: Path) -> None:
    source = tmp_path / "report.xlsx"; make_xlsx(source)
    task = await parse_upload(source, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert task["status"] == "succeeded"
    document = task["result"]["document"]
    assert document["document_type"] == "xlsx"
    assert document["tables"][0]["sheet_name"] == "Sales"


@pytest.mark.asyncio
async def test_real_pptx_parse_upload_task(tmp_path: Path) -> None:
    source = tmp_path / "presentation.pptx"; make_pptx(source)
    task = await parse_upload(source, "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert task["status"] == "succeeded"
    assert "Quarterly Results" in task["result"]["document"]["representations"]["plain_text"]


@pytest.mark.asyncio
async def test_legacy_xls_uses_automatic_pipeline(tmp_path: Path) -> None:
    soffice = shutil.which("soffice")
    if soffice is None: pytest.skip("当前环境未提供 LibreOffice")
    source_xlsx = tmp_path / "seed.xlsx"; legacy_dir = tmp_path / "legacy"; legacy_dir.mkdir(); make_xlsx(source_xlsx)
    created = subprocess.run([soffice, "--headless", "--convert-to", "xls", "--outdir", str(legacy_dir), str(source_xlsx)], capture_output=True, text=True, timeout=60, check=False)
    source = legacy_dir / "seed.xls"; assert created.returncode == 0 and source.is_file()
    task = await parse_upload(source, "application/vnd.ms-excel")
    assert task["status"] in {"succeeded", "partial"}
    assert task["plan"]["steps"][0]["skill_name"] == "office.parse_pipeline"
